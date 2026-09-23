import json
import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
import openpyxl

DEFAULT_SOURCE_FILES = {
    "water_bottles": r"D:\Mudra Branding Solutions\Corporate Mudhra\Water Bottles\WATER BOTTLES.xlsx",
    "pens": r"D:\Mudra Branding Solutions\Corporate Mudhra\Pens\XG - metal pen price list .xlsx",
    "electronics": r"D:\Mudra Branding Solutions\Corporate Mudhra\Electronics\ELECTRONICS AUG 2025.xlsx",
    "combos": r"D:\Mudra Branding Solutions\Corporate Mudhra\Combos\ALL combos price list .xlsx",
    "mugs": r"D:\Mudra Branding Solutions\Corporate Mudhra\Mugs\MUGS.xlsx",
    "notebooks": r"D:\Mudra Branding Solutions\Corporate Mudhra\Note Books\NOTEBOOK.xlsx",
}

CATEGORY_GST_RATES = {
    "Water Bottles": 12.0,
    "Writing Instruments": 18.0,
    "Electronics": 18.0,
    "Combos/Gift Sets": 18.0,
    "Mugs": 18.0,
    "Notebooks": 18.0,
}


def normalize_sku_key(sku: str) -> str:
    """
    Normalizes whitespace and casing for indexing/lookup without altering
    internal hyphens or silently merging different SKU formats.
    """
    if not sku:
        return ""
    # Collapse consecutive spaces into a single space and strip
    return re.sub(r"\s+", " ", str(sku)).strip().upper()


class PricingDataLoader:
    """
    Loads, normalizes, and validates pricing data from the six 2025 category
    workbooks under D:\\Mudra Branding Solutions\\Corporate Mudhra.
    """

    def __init__(self, source_files: Optional[Dict[str, str]] = None):
        self.source_files = source_files or DEFAULT_SOURCE_FILES
        self.pricing_version = "2025"
        self.records: List[Dict[str, Any]] = []
        self.report: Dict[str, Any] = {
            "generated_at": datetime.now().isoformat(),
            "pricing_version": self.pricing_version,
            "files_processed": [],
            "records_imported_per_category": {},
            "total_records_imported": 0,
            "records_skipped": 0,
            "missing_prices": [],
            "duplicate_skus": [],
            "quantity_tiers_discovered": {
                "Writing Instruments": ["1-299", "300-499", "500+"],
                "Water Bottles": ["1+ (open-ended)"],
                "Electronics": ["1+ (open-ended)"],
                "Combos/Gift Sets": ["1+ (open-ended)"],
                "Mugs": ["1+ (open-ended)"],
                "Notebooks": ["1+ (open-ended)"],
            },
            "gst_rates": CATEGORY_GST_RATES,
            "warnings": [],
            "errors": [],
        }

    def _create_record(
        self,
        sku: str,
        category: str,
        price: float,
        gst: float,
        qty_from: int,
        qty_to: Optional[int],
        source_file: str,
        source_sheet: str,
        source_row: int,
        source_category: str,
        timestamp: str,
    ) -> Dict[str, Any]:
        return {
            "sku": sku,
            "normalized_sku": normalize_sku_key(sku),
            "category": category,
            "unit_price_excl_gst": round(float(price), 2),
            "gst_percentage": float(gst),
            "quantity_from": int(qty_from),
            "quantity_to": int(qty_to) if qty_to is not None else None,
            "source_file": source_file,
            "source_sheet": source_sheet,
            "source_row": int(source_row),
            "source_category": source_category,
            "imported_at": timestamp,
            "pricing_version": self.pricing_version,
        }

    def load_water_bottles(self) -> List[Dict[str, Any]]:
        path = self.source_files["water_bottles"]
        fname = os.path.basename(path)
        cat = "Water Bottles"
        gst = CATEGORY_GST_RATES[cat]
        ts = datetime.now().isoformat()

        if not os.path.exists(path):
            self.report["errors"].append(f"File not found: {path}")
            return []

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"]
        records: List[Dict[str, Any]] = []
        seen: Dict[str, int] = {}

        for r in range(4, ws.max_row + 1):
            code_val = ws.cell(r, 2).value
            if not code_val:
                continue
            sku_str = str(code_val).strip()
            if "PRODUCT CODE" in sku_str.upper() or "SL. NO" in sku_str.upper():
                continue

            price_val = ws.cell(r, 4).value
            if price_val is None or not isinstance(price_val, (int, float)):
                self.report["missing_prices"].append({
                    "sku": sku_str,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "row": r,
                    "raw_value": str(price_val),
                })
                self.report["records_skipped"] += 1
                continue

            norm_key = normalize_sku_key(sku_str)
            if norm_key in seen:
                orig_row = seen[norm_key]
                self.report["duplicate_skus"].append({
                    "sku": sku_str,
                    "normalized_sku": norm_key,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "duplicate_row": r,
                    "original_row": orig_row,
                    "price": float(price_val),
                })
                self.report["warnings"].append(
                    f"Duplicate SKU {sku_str} at row {r} in {fname} (first seen row {orig_row}). Skipping row {r}."
                )
                self.report["records_skipped"] += 1
                continue

            seen[norm_key] = r
            rec = self._create_record(
                sku=sku_str,
                category=cat,
                price=float(price_val),
                gst=gst,
                qty_from=1,
                qty_to=None,
                source_file=fname,
                source_sheet="Sheet1",
                source_row=r,
                source_category=cat,
                timestamp=ts,
            )
            records.append(rec)

        wb.close()
        return records

    def load_pens(self) -> List[Dict[str, Any]]:
        path = self.source_files["pens"]
        fname = os.path.basename(path)
        cat = "Writing Instruments"
        gst = CATEGORY_GST_RATES[cat]
        ts = datetime.now().isoformat()

        if not os.path.exists(path):
            self.report["errors"].append(f"File not found: {path}")
            return []

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"]
        records: List[Dict[str, Any]] = []
        seen: Dict[str, int] = {}

        for r in range(4, ws.max_row + 1):
            code_val = ws.cell(r, 2).value
            if not code_val:
                continue
            sku_str = str(code_val).strip()
            if "PRODUCT CODE" in sku_str.upper() or "SL. NO" in sku_str.upper():
                continue

            # In XG - metal pen price list .xlsx:
            # Col 5 = MOQ 500
            # Col 6 = MOQ 300
            # Col 7 = MOQ 1-299
            p500 = ws.cell(r, 5).value
            p300 = ws.cell(r, 6).value
            p1_299 = ws.cell(r, 7).value

            if not (
                isinstance(p1_299, (int, float))
                and isinstance(p300, (int, float))
                and isinstance(p500, (int, float))
            ):
                self.report["missing_prices"].append({
                    "sku": sku_str,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "row": r,
                    "raw_p1_299": str(p1_299),
                    "raw_p300": str(p300),
                    "raw_p500": str(p500),
                })
                self.report["records_skipped"] += 1
                continue

            norm_key = normalize_sku_key(sku_str)
            if norm_key in seen:
                orig_row = seen[norm_key]
                self.report["duplicate_skus"].append({
                    "sku": sku_str,
                    "normalized_sku": norm_key,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "duplicate_row": r,
                    "original_row": orig_row,
                })
                self.report["warnings"].append(
                    f"Duplicate pen SKU {sku_str} at row {r} in {fname} (first seen row {orig_row}). Skipping variant row {r}."
                )
                self.report["records_skipped"] += 1
                continue

            seen[norm_key] = r

            # Tier 1: 1 - 299
            records.append(
                self._create_record(
                    sku=sku_str,
                    category=cat,
                    price=float(p1_299),
                    gst=gst,
                    qty_from=1,
                    qty_to=299,
                    source_file=fname,
                    source_sheet="Sheet1",
                    source_row=r,
                    source_category=cat,
                    timestamp=ts,
                )
            )
            # Tier 2: 300 - 499
            records.append(
                self._create_record(
                    sku=sku_str,
                    category=cat,
                    price=float(p300),
                    gst=gst,
                    qty_from=300,
                    qty_to=499,
                    source_file=fname,
                    source_sheet="Sheet1",
                    source_row=r,
                    source_category=cat,
                    timestamp=ts,
                )
            )
            # Tier 3: 500+
            records.append(
                self._create_record(
                    sku=sku_str,
                    category=cat,
                    price=float(p500),
                    gst=gst,
                    qty_from=500,
                    qty_to=None,
                    source_file=fname,
                    source_sheet="Sheet1",
                    source_row=r,
                    source_category=cat,
                    timestamp=ts,
                )
            )

        wb.close()
        return records

    def load_electronics(self) -> List[Dict[str, Any]]:
        path = self.source_files["electronics"]
        fname = os.path.basename(path)
        cat = "Electronics"
        gst = CATEGORY_GST_RATES[cat]
        ts = datetime.now().isoformat()

        if not os.path.exists(path):
            self.report["errors"].append(f"File not found: {path}")
            return []

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["ELECTRONICS"]
        records: List[Dict[str, Any]] = []
        seen: Dict[str, int] = {}

        for r in range(5, ws.max_row + 1):
            code_val = ws.cell(r, 2).value
            if not code_val:
                continue
            sku_str = str(code_val).strip()
            if "PRODUCT CODE" in sku_str.upper() or "SL. NO" in sku_str.upper():
                continue

            price_val = ws.cell(r, 4).value
            if price_val is None or not isinstance(price_val, (int, float)):
                self.report["missing_prices"].append({
                    "sku": sku_str,
                    "category": cat,
                    "file": fname,
                    "sheet": "ELECTRONICS",
                    "row": r,
                    "raw_value": str(price_val),
                })
                self.report["records_skipped"] += 1
                continue

            norm_key = normalize_sku_key(sku_str)
            if norm_key in seen:
                orig_row = seen[norm_key]
                self.report["duplicate_skus"].append({
                    "sku": sku_str,
                    "normalized_sku": norm_key,
                    "category": cat,
                    "file": fname,
                    "sheet": "ELECTRONICS",
                    "duplicate_row": r,
                    "original_row": orig_row,
                    "price": float(price_val),
                })
                self.report["warnings"].append(
                    f"Duplicate electronics SKU {sku_str} at row {r} in {fname} (first seen row {orig_row}). Skipping row {r}."
                )
                self.report["records_skipped"] += 1
                continue

            seen[norm_key] = r
            rec = self._create_record(
                sku=sku_str,
                category=cat,
                price=float(price_val),
                gst=gst,
                qty_from=1,
                qty_to=None,
                source_file=fname,
                source_sheet="ELECTRONICS",
                source_row=r,
                source_category=cat,
                timestamp=ts,
            )
            records.append(rec)

            # Generate canonical catalogue category alias (XG-T-xxx <-> XG-EL-xxx)
            t_match = re.match(r"^XG-T-(\d+)$", sku_str.upper())
            if t_match:
                el_sku = f"XG-EL-{t_match.group(1)}"
                el_rec = self._create_record(
                    sku=el_sku,
                    category=cat,
                    price=float(price_val),
                    gst=gst,
                    qty_from=1,
                    qty_to=None,
                    source_file=fname,
                    source_sheet="ELECTRONICS",
                    source_row=r,
                    source_category=cat,
                    timestamp=ts,
                )
                records.append(el_rec)

        wb.close()
        return records

    def load_mugs(self) -> List[Dict[str, Any]]:
        path = self.source_files["mugs"]
        fname = os.path.basename(path)
        cat = "Mugs"
        gst = CATEGORY_GST_RATES[cat]
        ts = datetime.now().isoformat()

        if not os.path.exists(path):
            self.report["errors"].append(f"File not found: {path}")
            return []

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"]
        records: List[Dict[str, Any]] = []
        seen: Dict[str, int] = {}

        # In MUGS.xlsx, Col 3 = PRODUCT CODE, Col 5 = PRICE
        for r in range(4, ws.max_row + 1):
            code_val = ws.cell(r, 3).value
            if not code_val:
                continue
            sku_str = str(code_val).strip()
            if "PRODUCT CODE" in sku_str.upper() or "SL. NO" in sku_str.upper():
                continue

            price_val = ws.cell(r, 5).value
            if price_val is None or not isinstance(price_val, (int, float)):
                self.report["missing_prices"].append({
                    "sku": sku_str,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "row": r,
                    "raw_value": str(price_val),
                })
                self.report["records_skipped"] += 1
                continue

            norm_key = normalize_sku_key(sku_str)
            if norm_key in seen:
                orig_row = seen[norm_key]
                self.report["duplicate_skus"].append({
                    "sku": sku_str,
                    "normalized_sku": norm_key,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "duplicate_row": r,
                    "original_row": orig_row,
                    "price": float(price_val),
                })
                self.report["warnings"].append(
                    f"Duplicate mug SKU {sku_str} at row {r} in {fname} (first seen row {orig_row}). Skipping row {r}."
                )
                self.report["records_skipped"] += 1
                continue

            seen[norm_key] = r
            rec = self._create_record(
                sku=sku_str,
                category=cat,
                price=float(price_val),
                gst=gst,
                qty_from=1,
                qty_to=None,
                source_file=fname,
                source_sheet="Sheet1",
                source_row=r,
                source_category=cat,
                timestamp=ts,
            )
            records.append(rec)

        wb.close()
        return records

    def load_notebooks(self) -> List[Dict[str, Any]]:
        path = self.source_files["notebooks"]
        fname = os.path.basename(path)
        cat = "Notebooks"
        gst = CATEGORY_GST_RATES[cat]
        ts = datetime.now().isoformat()

        if not os.path.exists(path):
            self.report["errors"].append(f"File not found: {path}")
            return []

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"]
        records: List[Dict[str, Any]] = []
        seen: Dict[str, int] = {}

        # In NOTEBOOK.xlsx, Col 2 = PRODUCT CODE, Col 4 = PRICE
        for r in range(4, ws.max_row + 1):
            code_val = ws.cell(r, 2).value
            if not code_val:
                continue
            sku_str = str(code_val).strip()
            if "PRODUCT CODE" in sku_str.upper() or "SL. NO" in sku_str.upper():
                continue

            price_val = ws.cell(r, 4).value
            if price_val is None or not isinstance(price_val, (int, float)):
                self.report["missing_prices"].append({
                    "sku": sku_str,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "row": r,
                    "raw_value": str(price_val),
                })
                self.report["warnings"].append(
                    f"Missing price for notebook SKU {sku_str} at row {r} in {fname}. Omitted from pricing master without guessing."
                )
                self.report["records_skipped"] += 1
                continue

            norm_key = normalize_sku_key(sku_str)
            if norm_key in seen:
                orig_row = seen[norm_key]
                self.report["duplicate_skus"].append({
                    "sku": sku_str,
                    "normalized_sku": norm_key,
                    "category": cat,
                    "file": fname,
                    "sheet": "Sheet1",
                    "duplicate_row": r,
                    "original_row": orig_row,
                    "price": float(price_val),
                })
                self.report["warnings"].append(
                    f"Duplicate notebook SKU {sku_str} at row {r} in {fname} (first seen row {orig_row}). Skipping row {r}."
                )
                self.report["records_skipped"] += 1
                continue

            seen[norm_key] = r
            rec = self._create_record(
                sku=sku_str,
                category=cat,
                price=float(price_val),
                gst=gst,
                qty_from=1,
                qty_to=None,
                source_file=fname,
                source_sheet="Sheet1",
                source_row=r,
                source_category=cat,
                timestamp=ts,
            )
            records.append(rec)

        wb.close()
        return records

    def load_combos(self) -> List[Dict[str, Any]]:
        path = self.source_files["combos"]
        fname = os.path.basename(path)
        cat = "Combos/Gift Sets"
        gst = CATEGORY_GST_RATES[cat]
        ts = datetime.now().isoformat()

        if not os.path.exists(path):
            self.report["errors"].append(f"File not found: {path}")
            return []

        wb = openpyxl.load_workbook(path, data_only=True)
        records: List[Dict[str, Any]] = []
        seen: Dict[str, Tuple[str, int]] = {}

        sheets = ["2 IN 1", "3 IN 1", "4 IN 1", "5 IN 1 ", "6 IN 1", "7 IN 1"]
        for sname in sheets:
            if sname not in wb.sheetnames:
                self.report["warnings"].append(f"Sheet {sname} not found in {fname}")
                continue
            ws = wb[sname]

            # Find all table header locations
            tables: List[Dict[str, int]] = []
            for r in range(1, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    val = ws.cell(r, c).value
                    if val and "PRODUCT CODE" in str(val).strip().upper():
                        tables.append({
                            "header_row": r,
                            "code_col": c,
                            "price_col": c + 2,
                        })

            for t in tables:
                h_row = t["header_row"]
                c_col = t["code_col"]
                p_col = t["price_col"]

                for r in range(h_row + 1, ws.max_row + 1):
                    code_val = ws.cell(r, c_col).value
                    if not code_val:
                        continue
                    sku_str = str(code_val).strip()
                    if "PRODUCT CODE" in sku_str.upper() or "SL. NO" in sku_str.upper():
                        break

                    price_val = ws.cell(r, p_col).value
                    if price_val is None or not isinstance(price_val, (int, float)):
                        self.report["missing_prices"].append({
                            "sku": sku_str,
                            "category": cat,
                            "file": fname,
                            "sheet": sname,
                            "row": r,
                            "raw_value": str(price_val),
                        })
                        self.report["records_skipped"] += 1
                        continue

                    norm_key = normalize_sku_key(sku_str)
                    if norm_key in seen:
                        orig_sheet, orig_row = seen[norm_key]
                        self.report["duplicate_skus"].append({
                            "sku": sku_str,
                            "normalized_sku": norm_key,
                            "category": cat,
                            "file": fname,
                            "sheet": sname,
                            "duplicate_row": r,
                            "original_sheet": orig_sheet,
                            "original_row": orig_row,
                            "price": float(price_val),
                        })
                        self.report["warnings"].append(
                            f"Duplicate combo SKU {sku_str} in sheet {sname} row {r} (first seen in {orig_sheet} row {orig_row}). Skipping."
                        )
                        self.report["records_skipped"] += 1
                        continue

                    seen[norm_key] = (sname, r)
                    rec = self._create_record(
                        sku=sku_str,
                        category=cat,
                        price=float(price_val),
                        gst=gst,
                        qty_from=1,
                        qty_to=None,
                        source_file=fname,
                        source_sheet=sname,
                        source_row=r,
                        source_category=f"Combos - {sname.strip()}",
                        timestamp=ts,
                    )
                    records.append(rec)

        wb.close()
        return records

    def load_all(self) -> List[Dict[str, Any]]:
        """Loads and combines all records from the six workbooks."""
        self.records = []
        self.report["files_processed"] = []

        loader_map = [
            ("water_bottles", "Water Bottles", self.load_water_bottles),
            ("pens", "Writing Instruments", self.load_pens),
            ("electronics", "Electronics", self.load_electronics),
            ("mugs", "Mugs", self.load_mugs),
            ("notebooks", "Notebooks", self.load_notebooks),
            ("combos", "Combos/Gift Sets", self.load_combos),
        ]

        for file_key, cat_name, load_fn in loader_map:
            fpath = self.source_files[file_key]
            self.report["files_processed"].append({
                "key": file_key,
                "category": cat_name,
                "path": fpath,
                "filename": os.path.basename(fpath),
            })
            cat_records = load_fn()
            self.report["records_imported_per_category"][cat_name] = len(cat_records)
            self.records.extend(cat_records)

        self.report["total_records_imported"] = len(self.records)
        return self.records

    def save_pricing_master(self, output_path: str) -> None:
        """Saves normalized records to JSON."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(self.records, f, indent=2, ensure_ascii=False)

    def save_report(self, output_path: str) -> None:
        """Saves import report to JSON."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(self.report, f, indent=2, ensure_ascii=False)


if __name__ == "__main__":
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    master_path = os.path.join(base_dir, "data", "pricing_master.json")
    report_path = os.path.join(base_dir, "data", "pricing_import_report.json")

    loader = PricingDataLoader()
    records = loader.load_all()
    loader.save_pricing_master(master_path)
    loader.save_report(report_path)

    print(f"Import complete! Successfully saved {len(records)} records to {master_path}")
    print(f"Audit report saved to {report_path}")
    print("Breakdown per category:")
    for cat, cnt in loader.report["records_imported_per_category"].items():
        print(f" - {cat}: {cnt}")
    print(f"Skipped records: {loader.report['records_skipped']}")
    print(f"Duplicate SKUs: {len(loader.report['duplicate_skus'])}")
    print(f"Missing prices: {len(loader.report['missing_prices'])}")
